import io
import csv
import re
import zipfile
from numbers import Number
from decimal import Decimal, InvalidOperation
import datetime
from typing import Tuple, Dict, Any, List, Optional
import numpy as np
import pandas as pd
from app.services.hs_product_service import (
    clean_text,
    clean_hs_code,
    get_hs8_product_info,
    get_product_name,
    get_product_names_bilingual,
    get_hs4_official_name,
    get_country_info,
    get_raw_processed_pair,
    RAW_VS_PROCESSED_PAIRS
)

MAX_IMPORT_ROWS = 200_000
MAX_EXPANDED_BYTES = 150 * 1024 * 1024
COLUMNS = [
    {'key': 'year', 'english': 'Year', 'khmer': 'ឆ្នាំ', 'required': True},
    {'key': 'month', 'english': 'Month', 'khmer': 'ខែ', 'required': True},
    {'key': 'hs8', 'english': 'HS8', 'khmer': 'កូដទំនិញ ៨ខ្ទង់', 'required': True},
    {'key': 'hs4', 'english': 'HS4', 'khmer': 'កូដទំនិញ ៤ខ្ទង់', 'required': True},
    {'key': 'desc', 'english': 'Description', 'khmer': 'បរិយាយកូដ ៤ខ្ទង់', 'required': False},
    {'key': 'country', 'english': 'Country', 'khmer': 'ប្រទេសគម្យដ្ឋាន', 'required': True},
    {'key': 'weight', 'english': 'Weight (kg)', 'khmer': 'ទម្ងន់ (គក្រ)', 'required': True},
    {'key': 'value', 'english': 'Value (KHR)', 'khmer': 'តម្លៃ (រៀល)', 'required': True},
]


def unit_price(value, weight):
    """Zero weight has no unit price; expose 0 consistently, never NaN/inf."""
    return value / weight if weight > 0 else 0.0


def label(value):
    return ' '.join(clean_text(value).casefold().split())


ALIASES = {label(c[n]): c['key'] for c in COLUMNS for n in ('english', 'khmer')}
ALIASES.update({label(k): v for k, v in {
    'year': 'year', 'month': 'month', 'hs8_code': 'hs8', 'hs4_code': 'hs4',
    'hs_8': 'hs8', 'hs_4': 'hs4', 'hs 8': 'hs8', 'hs 4': 'hs4',
    'hs8 code': 'hs8', 'hs4 code': 'hs4', 'hs_code': 'hs8', 'hs code': 'hs8',
    'country_code': 'country', 'country': 'country', 'destination': 'country',
    'weight_kg': 'weight', 'val_khr': 'value', 'value_khr': 'value',
    'weight kg': 'weight', 'value khr': 'value', 'weight': 'weight', 'value': 'value',
    'ប្រទេស': 'country', 'បរិយាយ': 'desc', 'បរិយាយកូដ': 'desc',
    'កូដ ៨ខ្ទង់': 'hs8', 'កូដ ៤ខ្ទង់': 'hs4', 'កូដ៨ខ្ទង់': 'hs8', 'កូដ៤ខ្ទង់': 'hs4',
    'ទម្ងន់': 'weight', 'តម្លៃ': 'value',
}.items()})


def month_bound(value):
    if not isinstance(value, str) or not re.fullmatch(r'(19\d{2}|20\d{2}|21\d{2}|2200)-(0[1-9]|1[0-2])', value):
        raise ValueError('Month bounds must be YYYY-MM (1900–2200)')
    return value


def decimal_number(value):
    if isinstance(value, bool):
        raise ValueError('Boolean values are not numbers')
    if isinstance(value, (Number, Decimal)):
        # XLSX stores IEEE-754 numbers with Excel's 15 significant digits.
        # Remove binary representation noise, not signs/currencies or real text precision.
        number = Decimal(format(value, '.15g') if isinstance(value, (float, np.floating)) else str(value))
        if not number.is_finite():
            raise ValueError('Nonfinite numbers are not permitted')
        text = format(number, 'f')
    else:
        text = clean_text(value)
    # Only conventional thousands separators; do not strip signs/currency/units.
    if ',' in text:
        if not re.fullmatch(r'\+?\d{1,3}(,\d{3})+(\.\d{1,6})?', text):
            raise ValueError('Malformed thousands separators')
        text = text.replace(',', '')
    if not re.fullmatch(r'\+?\d+(\.\d{1,6})?', text):
        raise ValueError('Expected a nonnegative number, at most 6 decimals, no currency/units')
    try:
        result = Decimal(text)
    except InvalidOperation:
        raise ValueError('Invalid number')
    if not result.is_finite() or result < 0 or result >= Decimal('1e18'):
        raise ValueError('Number outside permitted range')
    return format(result.normalize(), 'f')


def integer_number(value, low, high):
    result = Decimal(decimal_number(value))
    if result != result.to_integral_value() or not low <= result <= high:
        raise ValueError(f'Expected integer {low}–{high}')
    return int(result)


def hs_code(value, digits):
    text = clean_text(value)
    if re.fullmatch(r'\d+\.0+', text):
        text = text.split('.')[0]
    text = re.sub(r'[\s.\-_]', '', text)
    if not re.fullmatch(r'\d{1,' + str(digits) + r'}', text):
        raise ValueError(f'HS code must contain at most {digits} digits')
    result = text.zfill(digits)
    if int(result) == 0:
        raise ValueError('Zero HS code is not valid')
    return result


def parse_import(contents, filename, mode, start_month=None, end_month=None):
    """Bounded source parser with source-row diagnostics, never a totals merge."""
    result = {'rows': [], 'errors': [], 'duplicate_count': 0, 'column_mapping': {},
              'start_month': start_month, 'end_month': end_month}
    workbook = None
    try:
        if mode not in ('complete', 'partial'):
            raise ValueError('Mode must be complete or partial')
        if mode == 'complete' and not (start_month and end_month):
            raise ValueError('Complete imports require explicit start_month and end_month')
        if bool(start_month) != bool(end_month):
            raise ValueError('Supply both month bounds or neither')
        if start_month:
            month_bound(start_month)
            month_bound(end_month)
            if start_month > end_month:
                raise ValueError('start_month exceeds end_month')
        extension = filename.rsplit('.', 1)[-1].lower()
        if extension == 'csv':
            if b'\x00' in contents:
                raise ValueError('CSV must be UTF-8 text')
            iterator = csv.reader(io.StringIO(contents.decode('utf-8-sig')), strict=True)
        elif extension == 'xlsx':
            with zipfile.ZipFile(io.BytesIO(contents)) as archive:
                if len(archive.infolist()) > 1000 or sum(f.file_size for f in archive.infolist()) > MAX_EXPANDED_BYTES:
                    raise ValueError('Expanded workbook exceeds safe limits')
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=False, keep_links=False)
            if len(workbook.worksheets) != 1:
                raise ValueError('Import workbook must have exactly one worksheet')
            sheet = workbook.worksheets[0]
            if (sheet.max_column or 0) > 32 or (sheet.max_row or 0) > MAX_IMPORT_ROWS + 15:
                raise ValueError('Workbook dimensions exceed safe limits')
            iterator = sheet.iter_rows(values_only=True, max_col=32)
        else:
            raise ValueError('Only .csv and .xlsx are supported')
        header = None
        for row_num, cells in enumerate(iterator, 1):
            if row_num > MAX_IMPORT_ROWS + 15:
                raise ValueError('Import exceeds 200,000 rows')
            if len(cells) > 32 or any(len(str(v)) > 4096 for v in cells if v is not None):
                raise ValueError(f'Row {row_num}: too many columns or oversized cell')
            if not any(v is not None and clean_text(v) for v in cells):
                continue
            if header is None:
                recognized = {ALIASES[label(v)] for v in cells if label(v) in ALIASES}
                if not {'year', 'month', 'hs8', 'country'}.issubset(recognized):
                    if row_num < 15:
                        continue
                    raise ValueError('Required header not found in first 15 rows')
                mapping = {}
                for i, v in enumerate(cells):
                    key = ALIASES.get(label(v))
                    if key:
                        if key in mapping:
                            raise ValueError(f'Ambiguous duplicate column for {key}')
                        mapping[key] = i
                missing = [c['english'] for c in COLUMNS if c['required'] and c['key'] not in mapping]
                if missing:
                    raise ValueError('Missing required columns: ' + ', '.join(missing) + '. Weight must explicitly be kg and value KHR.')
                header = mapping
                result['column_mapping'] = {key: clean_text(cells[i]) for key, i in mapping.items()}
                continue
            try:
                values = {key: cells[i] if i < len(cells) else None for key, i in header.items()}
                if any(isinstance(v, str) and v.startswith('=') for v in values.values()):
                    raise ValueError('Formula cells are not permitted')
                year = integer_number(values['year'], 1900, 2200)
                month = integer_number(values['month'], 1, 12)
                period = f'{year:04d}-{month:02d}'
                if start_month and not start_month <= period <= end_month:
                    raise ValueError('Row is outside declared month coverage')
                code8, code4 = hs_code(values['hs8'], 8), hs_code(values['hs4'], 4)
                if code8[:4] != code4:
                    raise ValueError('HS8 does not belong to supplied HS4 parent')
                country = clean_text(values['country']) if values['country'] is not None else ''
                if not country or len(country) > 160:
                    raise ValueError('Country is required, at most 160 characters')
                result['rows'].append({'year': year, 'month': month, 'hs8_code': code8, 'hs4_code': code4,
                    'country_label': country, 'weight_kg': decimal_number(values['weight']),
                    'val_khr': decimal_number(values['value']), 'source_row': row_num})
            except (ValueError, TypeError) as exc:
                result['errors'].append({'row': row_num, 'message': str(exc)})
                if len(result['errors']) >= 500:
                    raise ValueError('Stopped after 500 source-row errors')
        if header is None:
            raise ValueError('Required trade schema header not found')
        if not result['rows']:
            raise ValueError('Empty imports cannot replace trade data; use a separate reviewed delete workflow')
        if not start_month:
            periods = [f"{r['year']:04d}-{r['month']:02d}" for r in result['rows']]
            result['start_month'], result['end_month'] = min(periods), max(periods)
    except Exception as exc:
        result['errors'].append({'row': 0, 'message': str(exc) if isinstance(exc, ValueError) else 'Could not parse UTF-8 CSV or XLSX file'})
    finally:
        if workbook:
            workbook.close()
    return result


class ETLService:
    @staticmethod
    def inspect_and_load_file(file_contents: bytes, filename: str) -> pd.DataFrame:
        """Compatibility parser; strict units/schema, never infer missing metrics."""
        if len(file_contents) > 26214400:
            raise ValueError('Import exceeds upload limit')
        parsed = parse_import(file_contents, filename, 'partial')
        if parsed['errors']:
            raise ValueError(str(parsed['errors']))
        return pd.DataFrame(parsed['rows']).rename(columns={'hs8_code': 'HS8', 'hs4_code': 'HS4',
               'country_label': 'Country', 'weight_kg': 'Weight (kg)', 'val_khr': 'Value (KHR)'})

    @staticmethod
    def resolve_period_overlaps(frames: List[pd.DataFrame]) -> pd.DataFrame:
        """Legacy implicit replacement is forbidden; use declared staged coverage."""
        if not frames:
            return pd.DataFrame()
        if len(frames) == 1:
            return frames[0]
        raise ValueError('Multiple files require explicit complete/partial coverage; row counts never imply completeness')

    @staticmethod
    def validate_trade_schema(df: pd.DataFrame) -> Tuple[bool, str, Dict[str, str]]:
        """Exact bilingual headers. kg/KHR are explicit, not guessed or converted."""
        col_map = {}
        for column in df.columns:
            key = ALIASES.get(label(column))
            if key:
                if key in col_map:
                    return False, f'Ambiguous duplicate column for {key}', col_map
                col_map[key] = column
        missing = [c['key'] for c in COLUMNS if c['required'] and c['key'] not in col_map]
        if missing:
            return False, 'Missing required kg/KHR trade columns: ' + ', '.join(missing), col_map
        return True, 'Dataset schema validated. Source rows require staged validation.', col_map

    @staticmethod
    def process_trade_data(df: pd.DataFrame, col_map: Dict[str, str] = None, exchange_rate: float = 4000.0, registry: Optional[Dict] = None) -> Tuple[pd.DataFrame, Dict[str, Any]]:
        """Compute metrics from an already scoped frame, using revision-pinned names.

        Resolvers are local to this invocation: no module dictionary mutations or
        cross-request naming state. Registry is mandatory for production callers.
        """
        if not np.isfinite(exchange_rate) or exchange_rate <= 0:
            raise ValueError('Exchange rate must be finite and positive')
        if registry is None:
            raise ValueError('A revision-pinned registry is required')
        df_clean = df.copy()
        countries = {entry['name_en']: entry for entry in registry['country'].values()}
        countries.update(registry['country'])

        def get_country_info(code):
            entry = countries[code]
            return {'EN': entry['name_en'], 'Canonical_KH': entry['name_km'], 'ISO': entry.get('iso3') or 'OTH'}

        def get_hs4_official_name(code):
            entry = registry['hs4'][code]
            return entry['name_en'], entry['name_km'], entry['name_en'], entry['name_km']

        def get_hs8_product_info(code):
            entry = registry['hs8'][code]
            parent = registry['hs4'][entry['parent_code']]
            return {'name_en': entry['name_en'], 'name_kh': entry['name_km'], 'category_en': parent['name_en'], 'category_kh': parent['name_km']}

        def get_product_name(code, raw_description=''):
            return registry['hs8'][code]['name_en']

        def get_product_names_bilingual(code):
            entry = registry['hs8'][code]
            return {'en': entry['name_en'], 'km': entry['name_km']}
        
        # Standardize column names if col_map provided
        if col_map:
            rename_dict = {v: k for k, v in col_map.items() if v is not None}
            df_clean = df_clean.rename(columns=rename_dict)
        
        # Clean text across all string columns
        for col in df_clean.select_dtypes(include=['object', 'string']).columns:
            df_clean[col] = df_clean[col].apply(clean_text)

        if not {'year', 'month', 'country'}.issubset(df_clean.columns):
            raise ValueError('Monthly grain and canonical country are required')
        df_clean['year'] = df_clean['year'].map(lambda value: integer_number(value, 1900, 2200))
        df_clean['month'] = df_clean['month'].map(lambda value: integer_number(value, 1, 12))
        
        # Standardize Country Attributes
        def resolve_country(raw_c):
            info = get_country_info(raw_c)
            return info['EN'], info['Canonical_KH'], info['ISO']

        country_data = df_clean['country'].apply(resolve_country)
        df_clean['country_en'] = country_data.apply(lambda x: x[0])
        df_clean['country_kh'] = country_data.apply(lambda x: x[1])
        df_clean['country_iso3'] = country_data.apply(lambda x: x[2])
        # Use clean English country name as default primary country column
        df_clean['country'] = df_clean['country_en']

        # Validate signs/units; never strip a minus sign or silently replace NaN.
        for target, source in (('weight_kg', 'weight'), ('val_khr', 'value')):
            if target not in df_clean.columns:
                if source not in df_clean.columns:
                    raise ValueError(f'Missing {target}')
                df_clean[target] = df_clean[source]
            df_clean[target] = df_clean[target].map(lambda value: float(decimal_number(value)))
        
        # Compute USD value dynamically based on exchange_rate input
        df_clean['val_usd'] = df_clean['val_khr'] / exchange_rate
        
        # Format HS codes with fixed widths
        if 'hs4_code' not in df_clean.columns:
            hs4_src = df_clean['hs4']
            df_clean['hs4_code'] = hs4_src.apply(lambda x: hs_code(x, 4))
        else:
            df_clean['hs4_code'] = df_clean['hs4_code'].apply(lambda x: hs_code(x, 4))
            
        if 'hs8_code' not in df_clean.columns:
            hs8_src = df_clean['hs8']
            df_clean['hs8_code'] = hs8_src.apply(lambda x: hs_code(x, 8))
        else:
            df_clean['hs8_code'] = df_clean['hs8_code'].apply(lambda x: hs_code(x, 8))
            
        if 'hs4_desc' not in df_clean.columns:
            desc_src = df_clean['desc'] if 'desc' in df_clean.columns else 'Agricultural Product'
            df_clean['hs4_desc'] = desc_src.astype(str).apply(clean_text)
            
        # Attach Human-Readable Product Name to dataframe
        df_clean['product_name'] = df_clean.apply(
            lambda r: get_product_name(r['hs8_code'], r['hs4_desc']), axis=1
        )
        
        # Summary Calculations
        total_weight_kg = df_clean['weight_kg'].sum()
        total_weight_ton = total_weight_kg / 1000.0
        total_val_khr = df_clean['val_khr'].sum()
        total_val_usd = df_clean['val_usd'].sum()
        overall_avg_price_usd_per_kg = unit_price(total_val_usd, total_weight_kg)
        
        # 1. HS8 Commodity Aggregation (Primary Level)
        hs8_agg = df_clean.groupby(['hs8_code', 'hs4_code']).agg(
            weight_kg=('weight_kg', 'sum'),
            val_khr=('val_khr', 'sum'),
            val_usd=('val_usd', 'sum')
        ).reset_index()
        hs8_agg['weight_ton'] = hs8_agg['weight_kg'] / 1000.0
        
        hs8_agg['revenue_share_pct'] = (hs8_agg['val_usd'] / max(0.001, total_val_usd)) * 100.0
        hs8_agg['volume_share_pct'] = (hs8_agg['weight_kg'] / (total_weight_kg if total_weight_kg > 0 else 1)) * 100.0
        hs8_agg['unit_price_usd_per_kg'] = np.where(hs8_agg['weight_kg'] > 0, hs8_agg['val_usd'] / hs8_agg['weight_kg'], 0.0)

        hs8_details = hs8_agg['hs8_code'].apply(get_hs8_product_info)
        hs8_agg['name_en'] = hs8_details.apply(lambda d: d['name_en'])
        hs8_agg['name_km'] = hs8_details.apply(lambda d: d['name_kh'])
        hs8_agg['category_en'] = hs8_details.apply(lambda d: d['category_en'])
        hs8_agg['category_km'] = hs8_details.apply(lambda d: d['category_kh'])
        hs8_agg = hs8_agg.sort_values('val_usd', ascending=False)
        
        # 2. HS4 Chapter Aggregation (Macro Reference)
        hs4_agg = df_clean.groupby(['hs4_code']).agg(
            weight_kg=('weight_kg', 'sum'),
            val_khr=('val_khr', 'sum'),
            val_usd=('val_usd', 'sum')
        ).reset_index()
        hs4_agg['weight_ton'] = hs4_agg['weight_kg'] / 1000.0
        
        hs4_agg['revenue_share_pct'] = (hs4_agg['val_usd'] / max(0.001, total_val_usd)) * 100.0
        hs4_agg['volume_share_pct'] = (hs4_agg['weight_kg'] / (total_weight_kg if total_weight_kg > 0 else 1)) * 100.0
        hs4_agg['unit_price_usd_per_kg'] = np.where(hs4_agg['weight_kg'] > 0, hs4_agg['val_usd'] / hs4_agg['weight_kg'], 0.0)
        
        hs4_titles_en, hs4_titles_km, hs4_descs_en, hs4_descs_km = [], [], [], []
        for _, row in hs4_agg.iterrows():
            t_en, t_km, d_en, d_km = get_hs4_official_name(row['hs4_code'])
            hs4_titles_en.append(t_en)
            hs4_titles_km.append(t_km)
            hs4_descs_en.append(d_en)
            hs4_descs_km.append(d_km)

        hs4_agg['hs4_official_title'] = hs4_titles_en
        hs4_agg['hs4_title_en'] = hs4_titles_en
        hs4_agg['hs4_title_km'] = hs4_titles_km
        hs4_agg['hs4_full_desc_en'] = hs4_descs_en
        hs4_agg['hs4_full_desc_km'] = hs4_descs_km
        hs4_agg = hs4_agg.sort_values('val_usd', ascending=False)

        # 3. Country & Market Concentration (HHI & Market Share Table)
        country_agg = df_clean.groupby(['country', 'country_en', 'country_kh', 'country_iso3']).agg(
            val_khr=('val_khr', 'sum'),
            val_usd=('val_usd', 'sum'),
            weight_kg=('weight_kg', 'sum')
        ).reset_index()
        country_agg['weight_ton'] = country_agg['weight_kg'] / 1000.0
        country_agg['revenue_share_pct'] = (country_agg['val_usd'] / max(0.001, total_val_usd)) * 100.0
        country_agg['volume_share_pct'] = (country_agg['weight_kg'] / (total_weight_kg if total_weight_kg > 0 else 1)) * 100.0
        country_agg['unit_price_usd_per_kg'] = np.where(country_agg['weight_kg'] > 0, country_agg['val_usd'] / country_agg['weight_kg'], 0.0)
        country_agg['share_pct'] = country_agg['revenue_share_pct']
        
        hhi_score = (country_agg['revenue_share_pct'] ** 2).sum()
        country_agg = country_agg.sort_values('val_usd', ascending=False)
        
        # 4. Monthly Seasonality (Temporal Dynamics) & Date Range
        df_clean['yyyy_mm'] = df_clean['year'].astype(str) + '-' + df_clean['month'].astype(str).str.zfill(2)
        monthly_agg = df_clean.groupby('yyyy_mm').agg(
            weight_kton=('weight_kg', lambda x: x.sum()/1000000.0),
            val_million_usd=('val_usd', lambda x: x.sum()/1000000.0),
            val_usd=('val_usd', 'sum'),
            val_khr=('val_khr', 'sum'),
            weight_kg=('weight_kg', 'sum')
        ).reset_index().sort_values('yyyy_mm')
        monthly_agg['weight_ton'] = monthly_agg['weight_kg'] / 1000.0
        monthly_agg['unit_price'] = np.where(monthly_agg['weight_kg'] > 0, monthly_agg['val_usd'] / monthly_agg['weight_kg'], 0.0)
        months_list = monthly_agg['yyyy_mm'].tolist()
        
        def format_ym(ym_str):
            try:
                y, m = ym_str.split('-')
                dt = datetime.date(int(y), int(m), 1)
                return dt.strftime('%b %Y')
            except Exception:
                return ym_str

        if months_list:
            begin_date = format_ym(months_list[0])
            end_date = format_ym(months_list[-1])
        else:
            begin_date = None
            end_date = None

        years_available = sorted(df_clean['year'].dropna().unique().astype(int).tolist()) if 'year' in df_clean.columns else []
        countries_available = sorted(df_clean['country'].dropna().unique().tolist()) if 'country' in df_clean.columns else []
        
        # Multi-dimensional aggregations for interactive cross-chart slicing
        hs4_multi = df_clean.groupby(['hs4_code', 'year', 'country']).agg(
            val_usd=('val_usd', 'sum'),
            val_khr=('val_khr', 'sum'),
            weight_kg=('weight_kg', 'sum')
        ).reset_index()
        hs4_multi['weight_ton'] = hs4_multi['weight_kg'] / 1000.0
        hs4_multi['hs4_official_title'] = hs4_multi['hs4_code'].map(lambda c: get_hs4_official_name(c)[0])
        hs4_multi['hs4_title_en'] = hs4_multi['hs4_code'].map(lambda c: get_hs4_official_name(c)[0])
        hs4_multi['hs4_title_km'] = hs4_multi['hs4_code'].map(lambda c: get_hs4_official_name(c)[1])

        country_multi = df_clean.groupby(['country', 'hs4_code', 'hs8_code', 'year']).agg(
            val_usd=('val_usd', 'sum'),
            val_khr=('val_khr', 'sum'),
            weight_kg=('weight_kg', 'sum')
        ).reset_index()
        country_multi['weight_ton'] = country_multi['weight_kg'] / 1000.0
        cm_cinfo = country_multi['country'].apply(resolve_country)
        country_multi['country_en'] = cm_cinfo.apply(lambda x: x[0])
        country_multi['country_kh'] = cm_cinfo.apply(lambda x: x[1])
        country_multi['country_iso3'] = cm_cinfo.apply(lambda x: x[2])
        country_multi['revenue_share_pct'] = (country_multi['val_usd'] / max(0.001, total_val_usd)) * 100.0
        country_multi['volume_share_pct'] = (country_multi['weight_kg'] / (total_weight_kg if total_weight_kg > 0 else 1)) * 100.0
        country_multi['unit_price_usd_per_kg'] = np.where(country_multi['weight_kg'] > 0, country_multi['val_usd'] / country_multi['weight_kg'], 0.0)

        monthly_multi = df_clean.groupby(['yyyy_mm', 'hs4_code', 'hs8_code', 'country', 'year']).agg(
            weight_kton=('weight_kg', lambda x: x.sum()/1000000.0),
            val_million_usd=('val_usd', lambda x: x.sum()/1000000.0),
            val_usd=('val_usd', 'sum'),
            val_khr=('val_khr', 'sum'),
            weight_kg=('weight_kg', 'sum')
        ).reset_index().sort_values('yyyy_mm')
        monthly_multi['weight_ton'] = monthly_multi['weight_kg'] / 1000.0
        
        commodity_multi = df_clean.groupby(['hs4_code', 'hs8_code', 'country', 'year']).agg(
            val_usd=('val_usd', 'sum'),
            val_khr=('val_khr', 'sum'),
            weight_kg=('weight_kg', 'sum')
        ).reset_index()
        commodity_multi['weight_ton'] = commodity_multi['weight_kg'] / 1000.0
        commodity_multi['unit_price_usd_per_kg'] = np.where(commodity_multi['weight_kg'] > 0, commodity_multi['val_usd'] / commodity_multi['weight_kg'], 0.0)
        
        cm_info = commodity_multi['hs8_code'].apply(get_hs8_product_info)
        commodity_multi['name_en'] = cm_info.apply(lambda d: d['name_en'])
        commodity_multi['name_km'] = cm_info.apply(lambda d: d['name_kh'])
        commodity_multi['category_en'] = cm_info.apply(lambda d: d['category_en'])
        commodity_multi['category_km'] = cm_info.apply(lambda d: d['category_kh'])
        commodity_multi['product_name'] = commodity_multi['name_en']
        commodity_multi['revenue_share_pct'] = (commodity_multi['val_usd'] / max(0.001, total_val_usd)) * 100.0
        commodity_multi['volume_share_pct'] = (commodity_multi['weight_ton'] / (total_weight_ton if total_weight_ton > 0 else 1)) * 100.0

        # 5. Per-HS8 Commodity Details (for /compare page)
        commodity_details: dict = {}
        for hs8_code, grp in df_clean.groupby('hs8_code'):
            hs8_code_str = clean_hs_code(hs8_code, 8)
            hs4_code_str = clean_hs_code(grp['hs4_code'].iloc[0], 4)
            names = get_product_names_bilingual(hs8_code_str)

            tot_weight_kg  = float(grp['weight_kg'].sum())
            tot_val_usd    = float(grp['val_usd'].sum())
            tot_val_khr    = float(grp['val_khr'].sum())
            avg_price_usd  = unit_price(tot_val_usd, tot_weight_kg)
            avg_price_khr  = unit_price(tot_val_khr, tot_weight_kg)
            num_markets    = int(grp['country'].nunique())

            monthly = grp.groupby('yyyy_mm').agg(
                val_usd=('val_usd', 'sum'),
                val_khr=('val_khr', 'sum'),
                weight_kg=('weight_kg', 'sum')
            ).reset_index().sort_values('yyyy_mm')
            monthly['price_usd_per_kg'] = np.where(monthly['weight_kg'] > 0, monthly['val_usd'] / monthly['weight_kg'], 0.0)
            monthly['price_khr_per_kg'] = np.where(monthly['weight_kg'] > 0, monthly['val_khr'] / monthly['weight_kg'], 0.0)
            monthly_trend = monthly.to_dict(orient='records')

            mkt = grp.groupby('country').agg(
                val_usd=('val_usd', 'sum'),
                val_khr=('val_khr', 'sum'),
                weight_kg=('weight_kg', 'sum')
            ).reset_index()
            mkt['share_pct'] = (mkt['val_usd'] / max(0.001, tot_val_usd)) * 100.0
            mkt = mkt.sort_values('val_usd', ascending=False)
            market_breakdown = mkt.to_dict(orient='records')

            mkt_yearly = grp.groupby(['country', 'year']).agg(
                val_usd=('val_usd', 'sum'),
                val_khr=('val_khr', 'sum'),
                weight_kg=('weight_kg', 'sum')
            ).reset_index()
            yearly_markets = mkt_yearly.to_dict(orient='records')

            raw_vs_processed = None
            pair_info = get_raw_processed_pair(hs4_code_str, hs8_code_str)
            if pair_info and pair_info['raw_hs8'] in registry['hs8'] and pair_info['processed_hs8'] in registry['hs8']:
                partner_hs8 = pair_info['processed_hs8'] if pair_info['is_raw'] else pair_info['raw_hs8']
                partner_grp = df_clean[df_clean['hs8_code'] == partner_hs8]
                partner_weight_kg = float(partner_grp['weight_kg'].sum())
                partner_val_usd   = float(partner_grp['val_usd'].sum())

                raw_hs8      = pair_info['raw_hs8']
                proc_hs8     = pair_info['processed_hs8']
                raw_grp      = df_clean[df_clean['hs8_code'] == raw_hs8]
                proc_grp     = df_clean[df_clean['hs8_code'] == proc_hs8]
                raw_names    = get_product_names_bilingual(raw_hs8)
                proc_names   = get_product_names_bilingual(proc_hs8)

                raw_vs_processed = {
                    'raw_hs8':             raw_hs8,
                    'processed_hs8':       proc_hs8,
                    'raw_name_en':         raw_names['en'],
                    'raw_name_km':         raw_names['km'],
                    'processed_name_en':   proc_names['en'],
                    'processed_name_km':   proc_names['km'],
                    'raw_val_usd':         round(float(raw_grp['val_usd'].sum()), 2),
                    'raw_weight_kg':       round(float(raw_grp['weight_kg'].sum()), 2),
                    'raw_price_usd_per_kg':  round(unit_price(float(raw_grp['val_usd'].sum()), float(raw_grp['weight_kg'].sum())), 4),
                    'raw_price_khr_per_kg':  round(unit_price(float(raw_grp['val_khr'].sum()), float(raw_grp['weight_kg'].sum())), 2),
                    'proc_val_usd':        round(float(proc_grp['val_usd'].sum()), 2),
                    'proc_weight_kg':      round(float(proc_grp['weight_kg'].sum()), 2),
                    'proc_price_usd_per_kg': round(unit_price(float(proc_grp['val_usd'].sum()), float(proc_grp['weight_kg'].sum())), 4),
                    'proc_price_khr_per_kg': round(unit_price(float(proc_grp['val_khr'].sum()), float(proc_grp['weight_kg'].sum())), 2),
                }

            commodity_details[hs8_code_str] = {
                'hs8_code':           hs8_code_str,
                'hs4_code':           hs4_code_str,
                'name_en':            names['en'],
                'name_km':            names['km'],
                'total_val_usd':      round(tot_val_usd, 2),
                'total_val_khr':      round(tot_val_khr, 2),
                'total_weight_kg':    round(tot_weight_kg, 2),
                'total_weight_ton':   round(tot_weight_kg / 1000.0, 4),
                'avg_price_usd_per_kg': round(avg_price_usd, 4),
                'avg_price_khr_per_kg': round(avg_price_khr, 2),
                'num_markets':        num_markets,
                'total_transactions': int(len(grp)),
                'monthly_trend':      monthly_trend,
                'market_breakdown':   market_breakdown,
                'yearly_markets':     yearly_markets,
                'raw_vs_processed':   raw_vs_processed,
            }

        # Flat commodity list for compare-page and overview dropdowns
        commodity_list = [
            {
                'hs8_code':   code,
                'hs4_code':   det['hs4_code'],
                'name_en':    det['name_en'],
                'name_km':    det['name_km'],
                'val_usd':    det['total_val_usd'],
                'weight_kg':  det['total_weight_kg'],
            }
            for code, det in sorted(commodity_details.items())
        ]

        metrics = {
            'begin_date': begin_date,
            'end_date': end_date,
            'available_years': years_available,
            'available_countries': countries_available,
            'total_weight_ton': float(total_weight_ton),
            'total_val_khr': float(total_val_khr),
            'total_val_usd': float(total_val_usd),
            'total_transactions': int(len(df_clean)),
            'overall_avg_price_usd_per_kg': float(overall_avg_price_usd_per_kg),
            'hhi_score': float(hhi_score),
            'is_high_risk': bool(hhi_score > 2500),
            'hs8_commodities': hs8_agg.to_dict(orient='records'),
            'hs4_summary': hs4_agg.to_dict(orient='records'),
            'hs8_summary': hs8_agg.to_dict(orient='records'),
            'country_summary': country_agg.to_dict(orient='records'),
            'monthly_summary': monthly_agg.to_dict(orient='records'),
            'hs4_multi': hs4_multi.to_dict(orient='records'),
            'country_multi': country_multi.to_dict(orient='records'),
            'monthly_multi': monthly_multi.to_dict(orient='records'),
            'commodity_multi': commodity_multi.to_dict(orient='records'),
            'commodity_details': commodity_details,
            'commodity_list': commodity_list,
        }
        
        return df_clean, metrics
